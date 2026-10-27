"""
Genera los entregables de seguimiento: gráficos burndown (PNG), plantilla de burndown, registro DORA y
plantilla de rendimiento (XLSX con fórmulas). Ejecutar desde la raíz: python3 metricas/generar_metricas.py
Fuentes: cronograma oficial del documento técnico (sec. 15.2) y consulta del tablero Jira SCRUM del 27-oct-2026.
"""
import json, os, datetime as dt
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

R = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
B = os.path.join(R, "metricas", "burndown"); D = os.path.join(R, "metricas", "dora"); P = os.path.join(R, "metricas", "rendimiento")
for d in (B, D, P): os.makedirs(d, exist_ok=True)
HOY = "2026-10-27"
SPRINTS = [("S1", "Sprint 1", "2026-10-06", "2026-10-19", 22), ("S2", "Sprint 2", "2026-10-20", "2026-11-02", 21), ("S3", "Sprint 3", "2026-11-03", "2026-11-16", 26)]
PROY = ("P", "Proyecto (Sprints 1 a 3)", "2026-10-06", "2026-11-16", 69)
REG = {"S1": [("2026-10-06", 0), ("2026-10-19", 0)], "S2": [("2026-10-20", 0), ("2026-10-26", 0), ("2026-10-27", 21)], "S3": [], "P": [("2026-10-06", 0), ("2026-10-19", 0), ("2026-10-26", 0), ("2026-10-27", 43)]}
FUENTE = "Jira SCRUM, 27-oct-2026: hasta esa fecha ninguna historia había cambiado de estado; tras actualizar el tablero, Sprint 1 (22 pts) y Sprint 2 (21 pts) en Finalizado"

def habiles(a, b):
    d, f, out = dt.date.fromisoformat(a), dt.date.fromisoformat(b), []
    while d <= f:
        if d.weekday() < 5: out.append(d)
        d += dt.timedelta(days=1)
    return out

# ---------------- PNG ----------------
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
def grafico(sid, nombre, ini, fin, pts, archivo):
    dias = habiles(ini, fin); n = len(dias) - 1; idx = {d: i for i, d in enumerate(dias)}
    ideal = [pts * (1 - i / n) for i in range(len(dias))]
    fig, ax = plt.subplots(figsize=(8.6, 4.3), dpi=160)
    ax.plot(range(len(dias)), ideal, ls="--", color="#6d7c79", lw=2, label=f"Línea ideal ({pts} → 0)")
    reg = REG[sid]
    if reg:
        xs = [idx[dt.date.fromisoformat(f)] for f, _ in reg]; ys = [pts - a for _, a in reg]
        ax.plot(xs, ys, color="#eb6834", lw=2.4, marker="o", ms=6, mec="white", mew=1.5, label="Trabajo restante real (registros de Jira)")
        ax.annotate(f"{ys[-1]} pts pendientes", (xs[-1], ys[-1]), textcoords="offset points", xytext=(0, -16 if ys[-1] > pts * 0.15 else 10), ha="center", fontsize=9, color="#142321")
        if sid == "S1":
            ax.text(0.98, 0.55, "Los 22 puntos se aceptaron el 27-oct,\n8 días después del cierre programado.", transform=ax.transAxes, ha="right", fontsize=8.8, color="#485956", bbox=dict(boxstyle="round,pad=0.4", fc="#fff4d9", ec="#f0d48f"))
    hoy = dt.date.fromisoformat(HOY)
    if dias[0] <= hoy <= dias[-1]:
        ax.axvline(idx[hoy], color="#b9c3c0", ls=":", lw=1.5); ax.text(idx[hoy] + 0.15, pts * 1.07, "consulta Jira 27-oct", fontsize=8.5, color="#485956", va="top")
    ax.set_ylim(0, pts * 1.1); ax.set_xlim(-0.4, n + 0.4)
    ax.set_ylabel("Puntos de historia pendientes"); ax.set_xlabel("Días hábiles")
    ax.set_title(f"Burndown · {nombre}\n{ini} a {fin} · {pts} puntos comprometidos", loc="left", fontsize=11.5, fontweight="bold")
    ax.grid(axis="y", color="#e3e8e6"); ax.legend(frameon=False, loc="lower left", fontsize=9)
    paso = max(1, len(dias) // 10); ticks = list(range(0, len(dias), paso))
    ax.set_xticks(ticks); ax.set_xticklabels([dias[i].strftime("%d-%b") for i in ticks], fontsize=8.5)
    if not reg:
        ax.text(0.62, 0.72, "Sin registros de avance real: el sprint no ha iniciado.\nSolo se muestra la línea ideal.", transform=ax.transAxes, ha="center", fontsize=9.5, color="#485956",
                bbox=dict(boxstyle="round,pad=0.5", fc="#f6f7f6", ec="#d9dfdd"))
    fig.tight_layout(); fig.savefig(os.path.join(B, archivo)); plt.close(fig)
grafico(*PROY, "burndown_proyecto.png")
for s in SPRINTS: grafico(*s, f"burndown_{s[0].lower()}.png")

# ---------------- XLSX helpers ----------------
F = "Arial"; th = Side(style="thin", color="B9C3C0")
HEAD = dict(font=Font(name=F, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="1D5C63"), alignment=Alignment(wrap_text=True, vertical="center"))
INPUT = PatternFill("solid", fgColor="FFFF00"); BLUE = Font(name=F, color="0000FF"); NORMAL = Font(name=F)
def head(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        for k, v in HEAD.items(): setattr(x, k, v)
        x.border = Border(top=th, bottom=th, left=th, right=th)
def fuente(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F: c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None, size=c.font.size if c.font else 11)
def anchos(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

# ---------------- Burndown XLSX ----------------
wb = Workbook(); ws = wb.active; ws.title = "Leyenda"
txt = ["Radar de permanencia · Seguimiento del burndown", "",
       "Cómo usar: en cada hoja de sprint, escriba en las celdas AMARILLAS los puntos aceptados ACUMULADOS en la fecha (solo historias que cumplen la Definition of Done y fueron revisadas por otro integrante).",
       "Deje en blanco los días sin registro: el trabajo restante real solo se calcula donde hay dato; nunca se inventa.",
       "Texto azul = dato ingresado. Texto negro = fórmula. No edite las columnas de fórmula.",
       "Trabajo restante = puntos comprometidos − puntos aceptados.  Línea ideal: lineal en días hábiles, de los puntos comprometidos a 0.",
       f"Datos precargados: {FUENTE}.", "Cronograma oficial: documento técnico, sección 15.2. Sprint 0 (21-sep a 5-oct) no tiene puntos. Backlog condicionado HU-05 y HU-10 (6 pts) no se compromete.",
       "Festivos de Colombia no descontados (2-nov y 16-nov de 2026 son festivos): ajústelos si el equipo no trabaja esos días."]
for i, t in enumerate(txt, 1): ws.cell(row=i, column=1, value=t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11)
ws.column_dimensions["A"].width = 150
# Resumen
rs = wb.create_sheet("Resumen")
head(rs, 1, ["Sprint", "Inicio", "Fin", "Puntos comprometidos", "Puntos aceptados (último registro)", "Puntos pendientes", "Variación vs. ideal (pts)", "Estado"])
for i, (sid, nom, ini, fin, pts) in enumerate(SPRINTS + [PROY], 2):
    hoja = nom.split(" (")[0].replace(" ", "_")
    rs.cell(row=i, column=1, value=nom); rs.cell(row=i, column=2, value=ini); rs.cell(row=i, column=3, value=fin)
    rs.cell(row=i, column=4, value=f"='{hoja}'!B2").font = Font(name=F, color="008000")
    rs.cell(row=i, column=5, value=f"=IF(COUNT('{hoja}'!C6:C40)=0,\"Sin registros\",LOOKUP(2,1/('{hoja}'!C6:C40<>\"\"),'{hoja}'!C6:C40))").font = Font(name=F, color="008000")
    rs.cell(row=i, column=6, value=f"=IF(ISNUMBER(E{i}),D{i}-E{i},\"—\")")
    rs.cell(row=i, column=7, value=f"=IF(ISNUMBER(E{i}),E{i}-(D{i}-LOOKUP(2,1/('{hoja}'!C6:C40<>\"\"),'{hoja}'!B6:B40)),\"—\")").font = Font(name=F, color="008000")
    rs.cell(row=i, column=8, value=f"=IF(ISNUMBER(G{i}),IF(G{i}<-1,\"Atrasado\",IF(G{i}>1,\"Adelantado\",\"Al día\")),\"Sin datos\")")
rs.cell(row=7, column=1, value="Backlog condicionado (HU-05, HU-10)"); rs.cell(row=7, column=4, value=6).font = BLUE; rs.cell(row=7, column=8, value="No comprometido")
rs.cell(row=9, column=1, value="Variación = aceptados − puntos que la línea ideal esperaba completados en la fecha del último registro. Negativo = atrasado.")
anchos(rs, [34, 12, 12, 14, 18, 14, 16, 14])
for sid, nom, ini, fin, pts in SPRINTS + [PROY]:
    hoja = nom.split(" (")[0].replace(" ", "_"); s = wb.create_sheet(hoja)
    s["A1"] = nom; s["A1"].font = Font(name=F, bold=True, size=13)
    s["A2"] = "Puntos comprometidos"; s["B2"] = pts; s["B2"].font = BLUE; s["C2"] = "Fuente: documento técnico, sec. 15.2"
    s["A3"] = f"Período: {ini} a {fin} (días hábiles)"
    dias = habiles(ini, fin); n = len(dias)
    head(s, 5, ["Fecha", "Línea ideal (pts pendientes)", "Puntos aceptados acumulados (ingresar)", "Trabajo restante real", "Fuente del registro"])
    reg = dict(REG[sid])
    for j, d in enumerate(dias):
        r = 6 + j
        s.cell(row=r, column=1, value=d).number_format = "yyyy-mm-dd"
        s.cell(row=r, column=2, value=f"=ROUND($B$2*(1-{j}/{n - 1}),2)")
        c = s.cell(row=r, column=3, value=reg.get(d.isoformat())); c.fill = INPUT; c.font = BLUE
        s.cell(row=r, column=4, value=f"=IF(C{r}=\"\",\"\",$B$2-C{r})")
        if d.isoformat() in reg: s.cell(row=r, column=5, value=FUENTE)
    s.cell(row=6 + n + 1, column=1, value="El gráfico de esta hoja muestra la línea ideal; la línea real con sus registros está en burndown_*.png y en el prototipo (Seguimiento › Burndown). Las celdas vacías de «Trabajo restante real» indican días sin registro.").font = Font(name=F, italic=True, color="485956")
    anchos(s, [13, 16, 22, 16, 60])
    ch = LineChart(); ch.title = f"Burndown · {nom}"; ch.y_axis.title = "Puntos pendientes"; ch.x_axis.title = "Fecha"; ch.height = 8; ch.width = 18
    ch.add_data(Reference(s, min_col=2, min_row=5, max_row=5 + n), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=6, max_row=5 + n))
    ch.series[0].graphicalProperties.line.dashStyle = "dash"; ch.series[0].graphicalProperties.line.solidFill = "6D7C79"
    ch.x_axis.number_format = "dd-mmm"; ch.x_axis.delete = False; ch.y_axis.delete = False
    s.add_chart(ch, "G5")
# Historias
hs = wb.create_sheet("Historias")
head(hs, 1, ["HU", "Historia", "Sprint (documento)", "Puntos", "Ítem Jira", "Sprint en Jira", "Estado en Jira (27-oct-2026, tras actualizar)", "Aceptada según DoD (Sí/No)"])
jira = {x["key"]: x for x in json.load(open(os.path.join(R, "evidencias", "jira", "jira_snapshot_2026-10-27_post.json")))}
HU = [("HU-01", "Información unificada y confiable", "Sprint 1", 8, "SCRUM-6"), ("HU-02", "Definición única de renuncia temprana", "Sprint 1", 3, "SCRUM-7"), ("HU-03", "Evolución de la rotación temprana", "Sprint 1", 5, "SCRUM-8"),
      ("HU-04", "Mes crítico de renuncia", "Sprint 1", 3, "SCRUM-53"), ("HU-06", "Filtrar el dashboard", "Sprint 1", 3, "SCRUM-22"), ("HU-07", "Retención por supervisor", "Sprint 2", 3, "SCRUM-10"),
      ("HU-08", "Pruebas técnicas y permanencia", "Sprint 2", 3, "SCRUM-9"), ("HU-09", "Perfiles con mayor rotación", "Sprint 2", 5, "SCRUM-47"), ("HU-11", "Resumen de causas y recomendaciones", "Sprint 2", 5, "SCRUM-23"),
      ("—", "Dataset del modelo (puntos derivados 21 − 16)", "Sprint 2", 5, None), ("HU-12", "Semáforo de riesgo del candidato", "Sprint 3", 8, "SCRUM-21"), ("HU-13", "Explicación del riesgo", "Sprint 3", 5, "SCRUM-52"),
      ("HU-14", "Confiabilidad de la herramienta", "Sprint 3", 5, "SCRUM-49"), ("HU-15", "Herramienta no discriminatoria", "Sprint 3", 5, "SCRUM-50"), ("HU-16", "Protección de datos personales", "Sprint 3", 3, "SCRUM-51"),
      ("HU-05", "Costo de la rotación temprana", "Backlog condicionado", 3, "SCRUM-46"), ("HU-10", "Motivos declarados de renuncia", "Backlog condicionado", 3, "SCRUM-48")]
for i, (h, n_, sp, pts, k) in enumerate(HU, 2):
    j = jira.get(k) if k else None
    k = k or "SCRUM-54"; j = jira.get(k)
    for c, v in enumerate([h, n_, sp, pts, k, ", ".join(j["sprint"]) if j else "—", j["status"] if j else "—", "Sí" if j and j["status"] == "Finalizado" else "No"], 1):
        x = hs.cell(row=i, column=c, value=v); x.font = BLUE if c in (4, 8) else NORMAL
    hs.cell(row=i, column=8).fill = INPUT
hs.cell(row=len(HU) + 3, column=3, value="Total comprometido Sprints 1-3"); hs.cell(row=len(HU) + 3, column=4, value=f'=SUMIFS(D2:D{len(HU)+1},C2:C{len(HU)+1},"Sprint*")')
hs.cell(row=len(HU) + 4, column=3, value="Aceptados (Sí)"); hs.cell(row=len(HU) + 4, column=4, value=f'=SUMIFS(D2:D{len(HU)+1},C2:C{len(HU)+1},"Sprint*",H2:H{len(HU)+1},"Sí")')
anchos(hs, [8, 44, 20, 8, 11, 32, 22, 18])
for w in wb.worksheets: fuente(w)
wb.save(os.path.join(B, "burndown_seguimiento.xlsx"))

# ---------------- DORA XLSX ----------------
wb = Workbook(); ws = wb.active; ws.title = "Leyenda"
for i, t in enumerate(["Radar de permanencia · Métricas DORA", "",
    "Registre en la hoja «Registro» cada cambio que llegue a un DESPLIEGUE: publicación de una versión utilizable por el cliente en su entorno de uso.",
    "No cuentan como despliegue: un commit, un merge, una entrega académica ni una demo local.",
    "Inicio del cambio = fecha y hora del primer commit del cambio. Despliegue = fecha y hora en que la versión queda disponible.",
    "Fallo = el despliegue requirió corrección, reversión o intervención urgente. Recuperación = momento en que se verificó que el servicio volvió a funcionar.",
    "Celdas AMARILLAS / texto azul = ingresar. Texto negro = fórmula. La hoja «Resultados» se recalcula sola.",
    "Estado al 27-oct-2026: sin entorno productivo ni despliegues registrados → «Sin despliegues medidos».",
    "La hoja «Ejemplo demostrativo» usa datos inventados solo para mostrar las fórmulas; no son resultados del proyecto."], 1):
    ws.cell(row=i, column=1, value=t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11)
ws.column_dimensions["A"].width = 140
def hoja_registro(wb, nombre, filas):
    s = wb.create_sheet(nombre)
    head(s, 1, ["ID", "Cambio / ítem Jira", "Inicio del cambio (primer commit)", "Despliegue", "Exitoso (Sí/No)", "Provocó fallo (Sí/No)", "Inicio del incidente", "Recuperación verificada", "Tiempo de entrega (h)", "Tiempo de recuperación (h)"])
    for r in range(2, 52):
        f = filas[r - 2] if r - 2 < len(filas) else None
        for c in range(1, 9):
            x = s.cell(row=r, column=c, value=(f[c - 1] if f else None)); x.fill = INPUT; x.font = BLUE
            if c in (3, 4, 7, 8): x.number_format = "yyyy-mm-dd hh:mm"
        s.cell(row=r, column=9, value=f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),E{r}="Sí"),(D{r}-C{r})*24,"")').number_format = "0.0"
        s.cell(row=r, column=10, value=f'=IF(AND(F{r}="Sí",ISNUMBER(G{r}),ISNUMBER(H{r})),(H{r}-G{r})*24,"")').number_format = "0.0"
    anchos(s, [6, 34, 20, 18, 10, 10, 18, 18, 14, 14]); return s
def hoja_resultados(wb, nombre, reg, ini, fin, demo=False):
    s = wb.create_sheet(nombre)
    s["A1"] = "Inicio del período"; s["B1"] = dt.datetime.fromisoformat(ini); s["A2"] = "Fin del período"; s["B2"] = dt.datetime.fromisoformat(fin)
    for c in ("B1", "B2"): s[c].number_format = "yyyy-mm-dd"; s[c].font = BLUE; s[c].fill = INPUT
    s["A3"] = "Semanas observadas"; s["B3"] = "=(B2-B1)/7"; s["B3"].number_format = "0.0"
    s["A4"] = "Despliegues registrados"; s["B4"] = f"=COUNT('{reg}'!D2:D51)"
    s["A5"] = "Despliegues exitosos"; s["B5"] = f"=COUNTIFS('{reg}'!D2:D51,\">0\",'{reg}'!E2:E51,\"Sí\")"
    s["A6"] = "Despliegues con fallo"; s["B6"] = f"=COUNTIFS('{reg}'!D2:D51,\">0\",'{reg}'!F2:F51,\"Sí\")"
    head(s, 8, ["Métrica", "Definición", "Fórmula", "Período observado", "Valor", "Unidad", "Fuente de evidencia", "Estado de medición", "Interpretación", "Acción de mejora"])
    sin = 'IF($B$4=0,"Sin despliegues medidos",'
    rows = [
        ("Frecuencia de despliegue", "Despliegues exitosos por período", "Despliegues exitosos ÷ semanas observadas", "=IF($B$4=0,\"—\",$B$5/$B$3)", "despliegues/semana", "Tags/releases de GitHub y registro de publicaciones", f'={sin}"Medido")', "Mayor frecuencia = entregas más pequeñas y menos riesgo", "Definir qué cuenta como despliegue y registrar cada publicación"),
        ("Tiempo de entrega de cambios", "Desde el primer commit hasta el despliegue", "Promedio(despliegue − primer commit)", f"=IF(COUNT('{reg}'!I2:I51)=0,\"—\",AVERAGE('{reg}'!I2:I51))", "horas", "Historial de Git + registro de despliegues", f'={sin}IF(COUNT(\'{reg}\'!I2:I51)=0,"Datos insuficientes","Medido"))', "Menor = el valor llega antes al cliente", "Vincular cada ítem de Jira con su rama o pull request"),
        ("Tasa de fallos de cambios", "% de despliegues que requieren corrección, reversión o intervención urgente", "Despliegues con fallo ÷ total de despliegues × 100", "=IF($B$4=0,\"—\",$B$6/$B$4*100)", "%", "Registro de incidentes vinculado al despliegue", f'={sin}"Medido")', "Menor = cambios más confiables", "Registrar toda corrección urgente posterior a una publicación"),
        ("Tiempo de recuperación ante fallos", "Desde el inicio del incidente hasta la recuperación verificada", "Promedio(recuperación − inicio del incidente)", f"=IF(COUNT('{reg}'!J2:J51)=0,\"—\",AVERAGE('{reg}'!J2:J51))", "horas", "Registro de incidentes con hora de detección y verificación", f'={sin}IF(COUNT(\'{reg}\'!J2:J51)=0,"Sin fallos registrados","Medido"))', "Menor = el equipo restablece el servicio más rápido", "Anotar hora de detección y de verificación de cada incidente")]
    for i, rw in enumerate(rows, 9):
        vals = [rw[0], rw[1], rw[2], '=TEXT($B$1,"yyyy-mm-dd")&" a "&TEXT($B$2,"yyyy-mm-dd")', rw[3], rw[4], rw[5], rw[6], rw[7], rw[8]]
        for c, v in enumerate(vals, 1):
            x = s.cell(row=i, column=c, value=v); x.alignment = Alignment(wrap_text=True, vertical="top")
        s.cell(row=i, column=5).number_format = "0.00"
    if demo: s["A14"] = "EJEMPLO DEMOSTRATIVO: valores calculados con datos inventados. No son resultados del proyecto."; s["A14"].font = Font(name=F, bold=True, color="B42E2E")
    anchos(s, [30, 34, 30, 24, 10, 16, 32, 22, 30, 36])
hoja_registro(wb, "Registro", [])
hoja_resultados(wb, "Resultados", "Registro", "2026-09-21", "2026-11-16")
ej = [("E1", "Ejemplo: filtros del panorama", dt.datetime(2026, 10, 6, 9), dt.datetime(2026, 10, 8, 17), "Sí", "No", None, None),
      ("E2", "Ejemplo: gráfico por supervisor", dt.datetime(2026, 10, 13, 9), dt.datetime(2026, 10, 15, 12), "Sí", "Sí", dt.datetime(2026, 10, 15, 13), dt.datetime(2026, 10, 15, 16, 30)),
      ("E3", "Ejemplo: semáforo", dt.datetime(2026, 10, 20, 8), dt.datetime(2026, 10, 23, 10), "Sí", "No", None, None)]
hoja_registro(wb, "Ejemplo registro", ej)
hoja_resultados(wb, "Ejemplo demostrativo", "Ejemplo registro", "2026-10-06", "2026-10-26", demo=True)
for w in wb.worksheets: fuente(w)
wb.save(os.path.join(D, "dora_registro_y_resultados.xlsx"))

# ---------------- Rendimiento XLSX ----------------
rd = json.load(open(os.path.join(P, "resultados_rendimiento.json")))
wb = Workbook(); ws = wb.active; ws.title = "Procedimiento"
for i, t in enumerate(["Radar de permanencia · Indicadores de rendimiento del sistema", "",
    "Objetivos de diseño (RNF-07, propuesta del documento técnico; no son resultados): filtro < 3 s; semáforo < 2 s; referencia de volumen 50.000 registros.",
    "Procedimiento reproducible: 1) cd prototipo && npm install && npm run build   2) python3 pruebas/test_e2e.py",
    "La prueba abre dist/index.html en Chromium (Playwright), mide 10 cargas iniciales, 10 aperturas por pantalla, 20 cambios de filtro (acción → indicador actualizado),",
    "20 evaluaciones del semáforo (clic → resultado visible) y 10 recálculos completos sobre 50.000 registros remuestreados con semilla fija.",
    "Los tiempos de Playwright incluyen la sobrecarga de la herramienta, por lo que son una cota superior del tiempo que percibe un usuario.",
    "Para medir en el equipo del cliente: abra la pestaña Seguimiento › Rendimiento del prototipo y use «Ejecutar medición»; copie los valores en la hoja «Plantilla».",
    "Pendiente: medición con el dashboard final (Power BI) y con datos reales del cliente."], 1):
    ws.cell(row=i, column=1, value=t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11)
ws.column_dimensions["A"].width = 150
s = wb.create_sheet("Resultados medidos")
s["A1"] = f"Fecha y hora: {rd['fecha']}"; s["A2"] = f"Entorno: {rd['entorno']}"; s["A3"] = "Tamaño del dataset demo: 3.242 registros (2.528 con ventana completa); volumen: 50.000 registros remuestreados"
head(s, 5, ["Prueba", "Registros", "Repeticiones", "Mínimo (ms)", "Promedio (ms)", "Mediana (ms)", "Máximo (ms)", "Errores", "Umbral (ms)", "Resultado frente al umbral", "Evidencia"])
for i, f in enumerate(rd["filas"], 6):
    vals = [f["prueba"], f["registros"], f["n"], f["min"], f["promedio"], f["mediana"], f["max"], f["errores"], f["umbral"]]
    for c, v in enumerate(vals, 1):
        x = s.cell(row=i, column=c, value=v)
        if c in (4, 5, 6, 7): x.number_format = "0.0"
    s.cell(row=i, column=10, value=f'=IF(I{i}="","Informativo (sin umbral)",IF(AND(G{i}<I{i},H{i}=0),"Cumple","No cumple"))')
    s.cell(row=i, column=11, value="metricas/rendimiento/resultados_rendimiento.json; evidencias/capturas/PF-25_rendimiento.png")
anchos(s, [62, 11, 12, 12, 13, 13, 12, 9, 12, 22, 60])
t = wb.create_sheet("Plantilla")
head(t, 1, ["Fecha y hora", "Entorno (equipo, navegador)", "Prueba", "Tamaño del dataset", "Repeticiones", "Mínimo (ms)", "Promedio (ms)", "Mediana (ms)", "Máximo (ms)", "Errores", "Umbral (ms)", "Resultado", "Evidencia"])
ejemplo = ["2026-11-10 10:00", "EJEMPLO · Portátil Windows 11, Chrome", "Respuesta a un filtro", 3242, 20, 12.0, 18.0, 17.5, 25.0, 0, 3000]
for r in range(2, 22):
    for c in range(1, 12):
        x = t.cell(row=r, column=c, value=(ejemplo[c - 1] if r == 2 else None)); x.fill = INPUT; x.font = BLUE
    t.cell(row=r, column=12, value=f'=IF(I{r}="","",IF(K{r}="","Informativo",IF(AND(I{r}<K{r},J{r}=0),"Cumple","No cumple")))')
t["A23"] = "La fila 2 es un ejemplo de formato (valores inventados); reemplácela con la primera medición real."
anchos(t, [18, 34, 30, 12, 12, 12, 12, 12, 12, 9, 12, 14, 40])
for w in wb.worksheets: fuente(w)
wb.save(os.path.join(P, "rendimiento_resultados_y_plantilla.xlsx"))
print("ok")
